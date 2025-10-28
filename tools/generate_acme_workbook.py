#!/usr/bin/env python3
import os
import re
import sys
import random
from datetime import date, timedelta
from typing import List, Dict, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(ROOT, "save6", "SAFe6OrganizationStructure.md")
OUT_XLSX = os.path.join(ROOT, "ACME_Project_Team_Workload.xlsx")

random.seed(42)

# === Configurable planning parameters ===
# Total number of sprints to plan
SPRINT_COUNT = 8
# Calendar length of a sprint in days (e.g., 14 for a two-week sprint)
SPRINT_CALENDAR_LENGTH_DAYS = 14
# Working days used for capacity per sprint (e.g., 10 weekdays in a 2-week sprint)
WORKING_DAYS_PER_SPRINT = 10
# Hours per work day used for capacity calculation
HOURS_PER_WORK_DAY = 6.5
# Default capacity hours per sprint for each resource
DEFAULT_CAPACITY_HOURS_PER_SPRINT = WORKING_DAYS_PER_SPRINT * HOURS_PER_WORK_DAY

# Simple name pools per "universe" to assign themed names
MARVEL = [
    "Peter Parker","Steve Rogers","Tony Stark","Bruce Banner","Wanda Maximoff","Stephen Strange",
    "Scott Lang","T'Challa","Sam Wilson","Bucky Barnes","Carol Danvers","Clint Barton",
    "Gamora","Drax","Groot","Nebula","Mantis","Okoye","Shuri","Pepper Potts",
    "Johnny Storm","Ben Grimm","Kate Bishop","Monica Rambeau","Kamala Khan","Jennifer Walters",
    "Shang-Chi","Dane Whitman","Yelena Belova","Maria Hill","Phil Coulson","Jessica Jones",
    "Luke Cage","Matt Murdock","Danny Rand","Frank Castle","Quake Johnson","Hank Pym"
]
DC = [
    "John Stewart","Oliver Queen","Dinah Lance","Billy Batson","Zatanna Zatara","Tim Drake",
    "Jason Todd","Barbara Gordon","Cassandra Cain","Stephanie Brown","Roy Harper","Wally West",
    "Donna Troy","Kara Zor-El","Kaldur'ahm","Victor Zsasz","Rene Ramirez","Helena Bertinelli",
    "Rory Regan","Ryan Choi","Ted Kord","Jaime Reyes","Mari McCabe","Carter Hall","Kendra Saunders",
    "Eobard Thawne","Mera Nereus","Ravager Wilson","Talia al Ghul","Huntress Wayne"
]
BOND = [
    "James Bond","Felix Leiter","Alec Trevelyan","Vesper Lynd","Severine Moreau","Gareth Mallory",
    "Bill Tanner","Le Chiffre","Raoul Silva","Eve Moneypenny (Alt)","Q (Alt)","Camille Montes"
]
OLYMPIANS = [
    "Hera","Apollo","Artemis","Ares","Hephaestus","Demeter","Poseidon","Hestia","Dionysus"
]

UNIVERSE_POOLS = {
    "avengers": MARVEL,
    "guardians": MARVEL,
    "x-men": MARVEL,
    "fantastic four": MARVEL,
    "s.h.i.e.l.d": MARVEL,
    "justice league": DC,
    "teen titans": DC,
    "watchmen": DC,
    "olympus": OLYMPIANS,
    "portfolio": OLYMPIANS,
    "bond": BOND,
    "james bond": BOND,
}

used_names = set()

def pool_for_unit(unit: str) -> List[str]:
    u = (unit or "").lower()
    for key, pool in UNIVERSE_POOLS.items():
        if key in u:
            return pool
    # Fallback by ART keywords
    if any(k in u for k in ["art", "agile release train", "team"]):
        return MARVEL + DC
    return MARVEL + DC + BOND + OLYMPIANS


def themed_random_name(unit: str) -> str:
    pool = pool_for_unit(unit)
    # Try to find an unused name
    random.shuffle(pool)
    for name in pool:
        if name not in used_names:
            used_names.add(name)
            return name
    # If exhausted, synthesize a name with an index
    idx = len(used_names) + 1
    synth = f"{unit.title()} Member {idx}"
    used_names.add(synth)
    return synth


class Resource:
    def __init__(self, name: str, role: str, level: str, art: Optional[str], art_name: Optional[str], team: Optional[str], unit: Optional[str]):
        self.name = name
        self.role = role
        self.level = level
        self.art = art
        self.art_name = art_name
        self.team = team
        self.unit = unit

    def as_row(self):
        return [self.name, self.role, self.level, self.art or "", self.team or "", self.unit or ""]


def parse_org_markdown(md_text: str) -> Dict[str, List[Resource]]:
    resources: List[Resource] = []
    stakeholders: List[Resource] = []

    level = None  # Portfolio, Large Solution, ART
    current_art = None
    current_art_name = None
    current_team = None
    current_team_display = None
    team_specialties: Dict[str, List[str]] = {}

    # Regex for person-role lines like **Batman** - *Product Manager*
    person_role_re = re.compile(r"^\*\*(.+?)\*\*\s*-\s*\*(.+?)\*")
    art_header_re = re.compile(r"^###\s*\*\*(ART\s*\d+):\s*(.+?)\*\*")
    team_header_re = re.compile(r"^####\s*\*\*(Team\s*(Alpha|Beta)):\s*(.+?)\*\*")

    lines = md_text.splitlines()
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## **Portfolio Level"):
            level = "Portfolio"
            current_art = None
            current_team = None
            continue
        if line.startswith("## **Large Solution Level"):
            level = "Large Solution"
            current_art = None
            current_team = None
            continue
        m_art = art_header_re.match(line)
        if m_art:
            level = "ART"
            current_art = m_art.group(1)
            current_art_name = m_art.group(2)
            current_team = None
            continue
        m_team = team_header_re.match(line)
        if m_team:
            current_team = m_team.group(1)
            current_team_display = m_team.group(3)
            # initialize specialties store
            team_specialties.setdefault(current_team_display, [])
            continue
        # Capture specialties lines: - **Team**: ...
        if line.startswith("- **Team**:"):
            # e.g., "- **Team**: UI/UX developers, Mission Planning specialists"
            payload = line.split(":", 1)[1].strip()
            parts = [p.strip() for p in payload.split(",") if p.strip()]
            # normalize role names to singular/title case
            norm = []
            for p in parts:
                p = re.sub(r"\s+developers?", " Developer", p, flags=re.I)
                p = re.sub(r"\s+engineers?", " Engineer", p, flags=re.I)
                p = re.sub(r"\s+specialists?", " Specialist", p, flags=re.I)
                p = re.sub(r"\s+systems?", " System", p, flags=re.I)
                p = re.sub(r"\s+analysis", " Analyst", p, flags=re.I)
                norm.append(p.title())
            if current_team_display:
                team_specialties[current_team_display].extend(norm)
            continue

        # Person/Role lines
        m_pr = person_role_re.match(line)
        if m_pr:
            person = m_pr.group(1).strip()
            role = m_pr.group(2).strip()
            unit_hint = current_team_display or current_art_name or level
            # Record the resource
            res = Resource(person, role, level or "", current_art, current_art_name, current_team, unit_hint)
            resources.append(res)
            # Consider most leadership roles as stakeholders
            if role.lower() in {"portfolio epic owner","portfolio architect","portfolio coordinator","solution train engineer","solution architect","solution management","product manager","release train engineer","system architect","scrum master"}:
                stakeholders.append(res)
            continue

    # Generate additional team members based on specialties
    for team_name, roles in team_specialties.items():
        # Determine ART and level for the team by scanning resources already tied to this team_name
        team_context = next((r for r in resources if r.unit == team_name), None)
        level_ctx = team_context.level if team_context else "ART"
        art_ctx = team_context.art if team_context else None
        art_name_ctx = team_context.art_name if team_context else None
        team_label = team_context.team if team_context else None
        for role in roles:
            # Create 2 members for each specialty role
            for _ in range(2):
                name = themed_random_name(team_name)
                resources.append(Resource(name, role, level_ctx, art_ctx, art_name_ctx, team_label, team_name))

    return {"resources": resources, "stakeholders": stakeholders}


def next_monday(from_date: date) -> date:
    days_ahead = (0 - from_date.weekday() + 7) % 7  # Monday=0
    if days_ahead == 0:
        days_ahead = 7
    return from_date + timedelta(days=days_ahead)


def style_header(ws, row=1):
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFE599")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[row]:
        cell.font = bold
        cell.fill = fill
        cell.alignment = center
        cell.border = border


def set_col_widths(ws, widths: Dict[int, float]):
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_workbook(data: Dict[str, List[Resource]]):
    wb = Workbook()
    # Remove default sheet
    default = wb.active
    wb.remove(default)

    resources: List[Resource] = data["resources"]
    stakeholders: List[Resource] = data["stakeholders"]

    # Project Info
    ws = wb.create_sheet("Project_Info")
    ws.append(["Field", "Value"]) 
    ws.append(["Project Name", "ACME Systems Engineering - Drone Swarm System (Example)"])
    ws.append(["Description", "Workbook for team organization, stakeholders, and resource workload tracking."])
    ws.append(["Owner", "Moneypenny (Solution Management)"])
    ws.append(["Version", "1.0"])
    ws.append(["Last Generated", date.today().isoformat()])
    style_header(ws, 1)
    set_col_widths(ws, {1: 30, 2: 100})

    # Sprints
    ws_sprints = wb.create_sheet("Sprints")
    ws_sprints.append(["Sprint", "Start", "End", "Length (days)"])
    start = next_monday(date.today())
    for i in range(1, SPRINT_COUNT + 1):
        s = start + timedelta(days=(i-1)*SPRINT_CALENDAR_LENGTH_DAYS)
        e = s + timedelta(days=SPRINT_CALENDAR_LENGTH_DAYS-1)
        ws_sprints.append([f"Sprint {i}", s.isoformat(), e.isoformat(), SPRINT_CALENDAR_LENGTH_DAYS])
    style_header(ws_sprints, 1)
    set_col_widths(ws_sprints, {1: 12, 2: 14, 3: 14, 4: 14})

    # Stakeholders
    ws_st = wb.create_sheet("Stakeholders")
    ws_st.append(["Name","Role","Level","ART","Team","Unit","Influence","Interest","Engagement Plan","Contact"])
    for r in stakeholders:
        influence = "High" if r.level in ("Portfolio","Large Solution") or r.role.lower() in ("product manager","solution train engineer","solution architect") else "Medium"
        interest = "High" if r.role.lower() in ("product manager","release train engineer","system architect","scrum master") else "Medium"
        ws_st.append([r.name, r.role, r.level, r.art or "", r.team or "", r.unit or "", influence, interest, "", ""])
    style_header(ws_st, 1)
    set_col_widths(ws_st, {1: 22,2:22,3:14,4:10,5:14,6:24,7:12,8:12,9:30,10:24})

    # Organization
    ws_org = wb.create_sheet("Organization")
    ws_org.append(["Level","ART","ART Name","Team","Unit","Role","Name","Notes"])
    for r in resources:
        ws_org.append([r.level, r.art or "", r.art_name or "", r.team or "", r.unit or "", r.role, r.name, ""])
    style_header(ws_org, 1)
    set_col_widths(ws_org, {1:16,2:8,3:26,4:14,5:24,6:26,7:22,8:30})

    # Teams summary
    ws_tm = wb.create_sheet("Teams")
    ws_tm.append(["ART","Team","Unit","Role","Count"])
    # Simple aggregation: count roles per team
    from collections import defaultdict
    counts = defaultdict(int)
    units = {}
    for r in resources:
        key = (r.art or "", r.team or "", r.unit or "", r.role)
        counts[key] += 1
        units[(r.art or "", r.team or "")] = r.unit or ""
    for (art, team, unit, role), count in sorted(counts.items()):
        ws_tm.append([art, team, unit, role, count])
    style_header(ws_tm, 1)
    set_col_widths(ws_tm, {1:8,2:16,3:24,4:28,5:10})

    # Activities (Backlog) – provide a template and a few sample rows from cross-team dependencies
    ws_act = wb.create_sheet("Activities")
    ws_act.append(["ID","Title","Team","ART","Assigned To","Role","Sprint","Estimate (h)","Status","Priority","Dependency"]) 
    samples = [
        ("A-001","GCS-Airborne interface contract","Justice League Command","ART 1","Batman","Product Manager","Sprint 1",6,"Planned","High","Professor X"),
        ("A-002","Secure comms requirement baseline","Teen Titans Communication","ART 3","Raven","Product Manager","Sprint 2",8,"Planned","High","Nick Fury"),
        ("A-003","Swarm/Flight behavior API","X-Men Coordination","ART 6","Jean Grey","Product Manager","Sprint 2",10,"Planned","High","Professor X"),
        ("A-004","Integration test plan v1","Avengers Assembly","ART 7","Thor","Product Manager","Sprint 1",12,"Planned","High","ALL"),
    ]
    for row in samples:
        ws_act.append(list(row))
    style_header(ws_act, 1)
    set_col_widths(ws_act, {1:10,2:40,3:28,4:10,5:22,6:20,7:12,8:14,9:14,10:12,11:30})

    # RACI matrix – template
    ws_raci = wb.create_sheet("RACI")
    ws_raci.append(["Activity","Zeus","Athena","Hermes","M","Q","Moneypenny","Product Managers","RTEs","System Architects","Scrum Masters"]) 
    activities = [
        "PI Planning (Portfolio)",
        "Solution PI Planning",
        "Architecture Sync",
        "Monthly Stakeholder Review",
        "Solution Demo",
        "Security Review",
    ]
    raci_rows = {
        "PI Planning (Portfolio)": ["A","C","I","R","C","I","C","I","I","I"],
        "Solution PI Planning":     ["I","C","I","A","C","R","C","R","C","I"],
        "Architecture Sync":        ["I","A","I","C","R","I","I","I","R","I"],
        "Monthly Stakeholder Review":["A","C","R","I","I","R","C","I","I","I"],
        "Solution Demo":            ["I","I","I","A","C","R","R","R","C","C"],
        "Security Review":          ["I","C","I","C","C","I","C","I","R","I"],
    }
    for act in activities:
        ws_raci.append([act]+raci_rows.get(act, ["I"]*9))
    style_header(ws_raci, 1)
    set_col_widths(ws_raci, {1:36,2:10,3:10,4:10,5:10,6:10,7:14,8:18,9:10,10:16,11:16})

    # Risks & Issues – template
    ws_ri = wb.create_sheet("Risks_Issues")
    ws_ri.append(["ID","Type","Title","Description","Owner","Impact","Probability","Mitigation","Status","Target Sprint"])
    style_header(ws_ri, 1)
    set_col_widths(ws_ri, {1:10,2:10,3:30,4:60,5:22,6:10,7:12,8:40,9:12,10:14})

    # Resource Workload sheet
    ws_rw = wb.create_sheet("Resource_Workload")
    # Header
    base_headers = ["Resource","Role","Level","ART","Team","Unit","Cap Hrs/Sprint"]
    alloc_headers = [f"S{i} %" for i in range(1, SPRINT_COUNT + 1)]
    hours_headers = [f"S{i} Hrs" for i in range(1, SPRINT_COUNT + 1)]
    tail_headers = ["Total Hrs","Avg Util %","Over/Under Hrs"]
    ws_rw.append(base_headers + alloc_headers + hours_headers + tail_headers)
    style_header(ws_rw, 1)

    # Populate rows
    start_row = 2
    for idx, r in enumerate(resources, start=start_row):
        # Build the row with dynamic number of sprint columns
        row = [r.name, r.role, r.level, r.art or "", r.team or "", r.unit or "", DEFAULT_CAPACITY_HOURS_PER_SPRINT]
        row += [0] * SPRINT_COUNT  # allocations
        row += [0] * SPRINT_COUNT  # hours
        row += [0, 0, 0]           # totals
        ws_rw.append(row)
        # Index helpers
        cap_col = 7
        alloc_start_col = 8
        alloc_end_col = alloc_start_col + SPRINT_COUNT - 1
        hours_start_col = alloc_end_col + 1
        hours_end_col = hours_start_col + SPRINT_COUNT - 1
        total_col = hours_end_col + 1
        avg_col = total_col + 1
        over_col = avg_col + 1
        # Formulas for hours per sprint: hours = cap * percent
        for s in range(SPRINT_COUNT):
            alloc_col = alloc_start_col + s
            hrs_col = hours_start_col + s
            ws_rw.cell(row=idx, column=hrs_col).value = f"={get_column_letter(cap_col)}{idx}*{get_column_letter(alloc_col)}{idx}"
        # Total
        ws_rw.cell(row=idx, column=total_col).value = f"=SUM({get_column_letter(hours_start_col)}{idx}:{get_column_letter(hours_end_col)}{idx})"
        # Avg Util
        ws_rw.cell(row=idx, column=avg_col).value = f"=AVERAGE({get_column_letter(alloc_start_col)}{idx}:{get_column_letter(alloc_end_col)}{idx})"
        # Over/Under relative to sprint count
        ws_rw.cell(row=idx, column=over_col).value = f"={get_column_letter(total_col)}{idx}-{get_column_letter(cap_col)}{idx}*{SPRINT_COUNT}"

    # Widths
    widths = {1:22,2:22,3:14,4:8,5:16,6:28,7:14}
    for c in range(SPRINT_COUNT):
        widths[8+c] = 8
    for c in range(SPRINT_COUNT):
        widths[8+SPRINT_COUNT+c] = 10
    tail_start = 7 + 2*SPRINT_COUNT + 1
    widths[tail_start] = 12
    widths[tail_start+1] = 12
    widths[tail_start+2] = 14
    set_col_widths(ws_rw, widths)
    ws_rw.freeze_panes = "H2"  # freeze before allocations

    # Summary by ART (pivot-style via formulas)
    ws_art = wb.create_sheet("Summary_ART")
    # Headers
    art_headers = [
        "ART","Headcount","Total Cap/Sprint",
    ] + [f"S{i} Hrs" for i in range(1, SPRINT_COUNT + 1)] + [
        "Total Hrs","Overall Util %","Avg Util %","Over/Under Hrs"
    ]
    ws_art.append(art_headers)
    style_header(ws_art, 1)
    set_col_widths(ws_art, {1:10,2:12,3:16,4:10,5:10,6:10,7:10,8:10,9:10,10:10,11:10,12:12,13:16,14:12,15:16})

    # Distinct ARTs
    distinct_arts = sorted({(r.art or "").strip() for r in resources if (r.art or "").strip()})
    for i, art in enumerate(distinct_arts, start=2):
        ws_art.cell(row=i, column=1, value=art)
        a_ref = f"$A{i}"
        # Base counts and capacity (ART=D, CAP=G)
        ws_art.cell(row=i, column=2, value=f"=COUNTIF(Resource_Workload!$D:$D,{a_ref})")
        ws_art.cell(row=i, column=3, value=f"=SUMIF(Resource_Workload!$D:$D,{a_ref},Resource_Workload!$G:$G)")
        # Dynamic per-sprint hours from Resource_Workload
        alloc_start_col = 8
        hours_start_col = alloc_start_col + SPRINT_COUNT
        out_col = 4
        for s in range(SPRINT_COUNT):
            hrs_letter = get_column_letter(hours_start_col + s)
            ws_art.cell(row=i, column=out_col, value=f"=SUMIF(Resource_Workload!$D:$D,{a_ref},Resource_Workload!${hrs_letter}:${hrs_letter})")
            out_col += 1
        # Tail columns on Resource_Workload
        total_col_rw = hours_start_col + SPRINT_COUNT
        avg_col_rw = total_col_rw + 1
        over_col_rw = avg_col_rw + 1
        ws_art.cell(row=i, column=out_col, value=f"=SUMIF(Resource_Workload!$D:$D,{a_ref},Resource_Workload!${get_column_letter(total_col_rw)}:${get_column_letter(total_col_rw)})")
        out_col += 1
        ws_art.cell(row=i, column=out_col, value=f"=IFERROR({get_column_letter(out_col-1)}{i}/({get_column_letter(3)}{i}*{SPRINT_COUNT}),0)")
        out_col += 1
        ws_art.cell(row=i, column=out_col, value=f"=AVERAGEIF(Resource_Workload!$D:$D,{a_ref},Resource_Workload!${get_column_letter(avg_col_rw)}:${get_column_letter(avg_col_rw)})")
        out_col += 1
        ws_art.cell(row=i, column=out_col, value=f"=SUMIF(Resource_Workload!$D:$D,{a_ref},Resource_Workload!${get_column_letter(over_col_rw)}:${get_column_letter(over_col_rw)})")

    # Summary by Team (ART + Team + Unit)
    ws_team = wb.create_sheet("Summary_Team")
    team_headers = [
        "ART","Team","Unit","Headcount","Total Cap/Sprint",
    ] + [f"S{i} Hrs" for i in range(1, SPRINT_COUNT + 1)] + [
        "Total Hrs","Overall Util %","Avg Util %","Over/Under Hrs"
    ]
    ws_team.append(team_headers)
    style_header(ws_team, 1)
    set_col_widths(ws_team, {1:10,2:14,3:28,4:12,5:16,6:10,7:10,8:10,9:10,10:10,11:10,12:10,13:10,14:12,15:16,16:12,17:16})

    distinct_team_keys = sorted({(r.art or "", r.team or "", r.unit or "") for r in resources if (r.art or "").strip() and (r.team or "").strip()})
    for i, (art, team, unit) in enumerate(distinct_team_keys, start=2):
        ws_team.cell(row=i, column=1, value=art)
        ws_team.cell(row=i, column=2, value=team)
        ws_team.cell(row=i, column=3, value=unit)
        a_ref = f"$A{i}"  # ART
        b_ref = f"$B{i}"  # Team
        c_ref = f"$C{i}"  # Unit
        # COUNTIFS over ART(D), Team(E), Unit(F)
        ws_team.cell(row=i, column=4, value="=COUNTIFS(Resource_Workload!$D:$D,"+a_ref+",Resource_Workload!$E:$E,"+b_ref+",Resource_Workload!$F:$F,"+c_ref+")")
        ws_team.cell(row=i, column=5, value="=SUMIFS(Resource_Workload!$G:$G,Resource_Workload!$D:$D,"+a_ref+",Resource_Workload!$E:$E,"+b_ref+",Resource_Workload!$F:$F,"+c_ref+")")
        # Dynamic hours cols on Resource_Workload
        alloc_start_col = 8
        hours_start_col = alloc_start_col + SPRINT_COUNT
        out_col = 6
        for s in range(SPRINT_COUNT):
            hrs_letter = get_column_letter(hours_start_col + s)
            ws_team.cell(row=i, column=out_col, value=f"=SUMIFS(Resource_Workload!${hrs_letter}:${hrs_letter},Resource_Workload!$D:$D,{a_ref},Resource_Workload!$E:$E,{b_ref},Resource_Workload!$F:$F,{c_ref})")
            out_col += 1
        # Tail columns from RW
        total_col_rw = hours_start_col + SPRINT_COUNT
        avg_col_rw = total_col_rw + 1
        over_col_rw = avg_col_rw + 1
        ws_team.cell(row=i, column=out_col, value=f"=SUMIFS(Resource_Workload!${get_column_letter(total_col_rw)}:${get_column_letter(total_col_rw)},Resource_Workload!$D:$D,{a_ref},Resource_Workload!$E:$E,{b_ref},Resource_Workload!$F:$F,{c_ref})")
        out_col += 1
        ws_team.cell(row=i, column=out_col, value=f"=IFERROR({get_column_letter(out_col-1)}{i}/({get_column_letter(5)}{i}*{SPRINT_COUNT}),0)")
        out_col += 1
        ws_team.cell(row=i, column=out_col, value=f"=AVERAGEIFS(Resource_Workload!${get_column_letter(avg_col_rw)}:${get_column_letter(avg_col_rw)},Resource_Workload!$D:$D,{a_ref},Resource_Workload!$E:$E,{b_ref},Resource_Workload!$F:$F,{c_ref})")
        out_col += 1
        ws_team.cell(row=i, column=out_col, value=f"=SUMIFS(Resource_Workload!${get_column_letter(over_col_rw)}:${get_column_letter(over_col_rw)},Resource_Workload!$D:$D,{a_ref},Resource_Workload!$E:$E,{b_ref},Resource_Workload!$F:$F,{c_ref})")

    # Dashboard – lightweight KPIs
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.append(["Metric","Value"]) 
    ws_dash.append(["Total Resources", f"=COUNTA(Resource_Workload!A:A)-1"])
    # Determine the Avg Util % column letter dynamically
    alloc_start_col = 8
    hours_start_col = alloc_start_col + SPRINT_COUNT
    total_col_rw = hours_start_col + SPRINT_COUNT
    avg_col_rw = total_col_rw + 1
    avg_letter = get_column_letter(avg_col_rw)
    ws_dash.append(["Avg Utilization (all)", f"=AVERAGE(Resource_Workload!{avg_letter}:{avg_letter})"])
    # Note: embed comparison string carefully to avoid quote escaping issues
    ws_dash.append(["Overutilized Count (>1.0 avg)", f"=COUNTIF(Resource_Workload!{avg_letter}:{avg_letter},\">1\")"])
    style_header(ws_dash, 1)
    set_col_widths(ws_dash, {1:36,2:24})

    wb.save(OUT_XLSX)
    return OUT_XLSX


def main():
    if not os.path.exists(MD_PATH):
        print(f"Markdown not found: {MD_PATH}", file=sys.stderr)
        sys.exit(1)
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()
    data = parse_org_markdown(md_text)
    out = build_workbook(data)
    print(f"Workbook generated: {out}")


if __name__ == "__main__":
    main()
